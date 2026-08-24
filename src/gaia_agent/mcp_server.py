from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import json
import mimetypes
import os
import socket
import stat
import threading
import time
import zipfile
import zlib
from collections.abc import Awaitable, Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from mcp.server import MCPServer
from mcp.types import CallToolResult, TextContent
from openai import AsyncOpenAI
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf import filters as pypdf_filters
from pypdf.errors import LimitReachedError

from gaia_agent.models import HOSTED_USAGE_METADATA_KEY, Usage
from gaia_agent.providers.base import (
    SafeToolFailure,
    ToolCallError,
    serialize_tool_output,
    tool_output_size_bytes,
)
from gaia_agent.reliability import RetryPolicy, retry_delay
from gaia_agent.search import (
    OpenAISearchProvider,
    SearchProvider,
    SearchProviderChain,
    SearchProviderResult,
    SerpAPISearchProvider,
    SerperSearchProvider,
)

MAX_TOOL_OUTPUT_BYTES = 64 * 1024
MAX_WEB_SEARCH_OUTPUT_BYTES = 16 * 1024
MAX_WEB_SEARCH_SOURCES = 8
MAX_HOSTED_WEB_SEARCH_TIMEOUT_SECONDS = 60.0
MAX_DIRECT_WEBPAGE_TIMEOUT_SECONDS = 20.0
MAX_DOWNLOAD_BYTES = 8 * 1024 * 1024
MAX_ATTACHMENT_BYTES = 32 * 1024 * 1024
MAX_AUDIO_ATTACHMENT_BYTES = 25_000_000
MAX_REDIRECTS = 10
MAX_DECODED_TEXT_BYTES = 1024 * 1024
MAX_TABULAR_ROWS = 10_000
MAX_TABULAR_CELLS = 100_000
MAX_PDF_PAGES = 100
MAX_PDF_TEXT_BYTES = 1024 * 1024
MAX_PDF_DECODED_STREAM_BYTES = 2 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 1_000
MAX_XLSX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
MAX_XLSX_WORKSHEETS = 50
_SENSITIVE_KEY_PARTS = ("api_key", "authorization", "credential", "password", "secret", "token")
_AUDIO_TRANSCRIPTION_SUFFIXES = frozenset(
    {".flac", ".mp3", ".mp4", ".mpeg", ".mpga", ".m4a", ".ogg", ".wav", ".webm"}
)
_PDF_FILTER_LIMIT_NAMES = (
    "MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH",
    "JBIG2_MAX_OUTPUT_LENGTH",
    "LZW_MAX_OUTPUT_LENGTH",
    "RUN_LENGTH_MAX_OUTPUT_LENGTH",
    "ZLIB_MAX_OUTPUT_LENGTH",
)
_PDF_FILTER_LOCK = threading.Lock()
_TRACE_META_KEY = "gaia.trace"
_MAX_TRACE_COUNT = 1_000_000
_MAX_FALLBACK_URL_CHARS = 2_048


class _TrustedHostedResult(dict[str, Any]):
    pass


class _SearchFallbackResult(dict[str, Any]):
    def __init__(self, payload: dict[str, Any], trace_metadata: dict[str, Any]) -> None:
        super().__init__(payload)
        self.trace_metadata = trace_metadata


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True


def _replace_secrets(text: str, secrets: tuple[str, ...]) -> str:
    for secret in secrets:
        if secret:
            text = text.replace(secret, "[REDACTED]")
    return text


def _redact(value: Any, secrets: tuple[str, ...]) -> Any:
    if isinstance(value, dict):
        return {
            key: (
                "[REDACTED]"
                if any(part in str(key).lower() for part in _SENSITIVE_KEY_PARTS)
                else _redact(item, secrets)
            )
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_redact(item, secrets) for item in value]
    if isinstance(value, tuple):
        return tuple(_redact(item, secrets) for item in value)
    if isinstance(value, str):
        return _replace_secrets(value, secrets)
    return value


def _bounded_result(
    result: dict[str, Any],
    *,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    metadata = metadata or {}
    combined = {**result, **metadata}
    if tool_output_size_bytes(combined) <= MAX_TOOL_OUTPUT_BYTES:
        return combined
    payload = serialize_tool_output(result).encode("utf-8")
    low = 0
    high = min(len(payload), MAX_TOOL_OUTPUT_BYTES)
    best: dict[str, Any] = {"output": "", "truncated": True, **metadata}
    while low <= high:
        middle = (low + high) // 2
        candidate = {
            "output": payload[:middle].decode("utf-8", errors="ignore"),
            "truncated": True,
            **metadata,
        }
        candidate_size = tool_output_size_bytes(candidate)
        if candidate_size <= MAX_TOOL_OUTPUT_BYTES:
            best = candidate
            low = middle + 1
        else:
            high = middle - 1
    return best


def _bounded_web_search_result(
    result: dict[str, Any],
    *,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    metadata = metadata or {}
    payload = dict(result)
    answer = str(payload.get("answer") or "")
    raw_sources = payload.get("sources")
    sources = list(raw_sources[:MAX_WEB_SEARCH_SOURCES]) if isinstance(raw_sources, list) else []
    extra = {key: value for key, value in payload.items() if key not in {"answer", "sources"}}

    def candidate(answer_text: str) -> dict[str, Any]:
        return {"answer": answer_text, "sources": sources, **extra, **metadata}

    while sources and tool_output_size_bytes(candidate("")) > MAX_WEB_SEARCH_OUTPUT_BYTES:
        sources.pop()
    if tool_output_size_bytes(candidate("")) > MAX_WEB_SEARCH_OUTPUT_BYTES:
        extra.clear()
    full = candidate(answer)
    if tool_output_size_bytes(full) <= MAX_WEB_SEARCH_OUTPUT_BYTES:
        return full

    low = 0
    high = len(answer)
    best = candidate("")
    while low <= high:
        middle = (low + high) // 2
        bounded = candidate(answer[:middle])
        if tool_output_size_bytes(bounded) <= MAX_WEB_SEARCH_OUTPUT_BYTES:
            best = bounded
            low = middle + 1
        else:
            high = middle - 1
    return best


class WebReader(Protocol):
    async def read(self, url: str) -> dict[str, Any]: ...


class HostedBackend(Protocol):
    async def web_search(self, query: str) -> dict[str, Any]: ...

    async def execute_python(self, instruction: str, paths: list[str]) -> dict[str, Any]: ...

    async def transcribe_audio(self, path: str) -> dict[str, Any]: ...


class WebpageError(ToolCallError):
    def __init__(
        self,
        category: str,
        *,
        retryable: bool,
        http_status: int | None = None,
        attempt_count: int = 1,
    ) -> None:
        super().__init__(category)
        self.category = category
        self.retryable = retryable
        self.provider = None
        self.http_status = http_status
        self.attempt_count = attempt_count


class HTTPWebReader:
    def __init__(
        self,
        *,
        client: httpx.AsyncClient | None = None,
        timeout_seconds: float = 20,
        max_bytes: int = MAX_DOWNLOAD_BYTES,
    ) -> None:
        self._client = client or httpx.AsyncClient(follow_redirects=False, trust_env=False)
        self._owns_client = client is None
        self._timeout_seconds = timeout_seconds
        self._max_bytes = max_bytes

    async def read(self, url: str) -> dict[str, Any]:
        try:
            result = await self._read_once(url)
        except TimeoutError:
            raise WebpageError("timeout", retryable=True) from None
        except httpx.HTTPStatusError as error:
            status = error.response.status_code
            raise WebpageError(
                "http_status",
                retryable=status == 408 or status == 429 or status >= 500,
                http_status=status,
            ) from None
        except httpx.DecodingError:
            raise WebpageError("parse_failure", retryable=False) from None
        except httpx.TransportError:
            raise WebpageError("transport", retryable=True) from None
        if not result.get("text", "").strip():
            raise WebpageError("empty_content", retryable=False)
        return result

    async def _read_once(self, url: str) -> dict[str, Any]:
        async with asyncio.timeout(self._timeout_seconds):
            current_url = httpx.URL(url)
            for _ in range(MAX_REDIRECTS + 1):
                pinned_urls, headers, extensions = await _pin_public_url(current_url)
                redirect_url: httpx.URL | None = None
                last_connect_error: httpx.ConnectError | httpx.ConnectTimeout | None = None
                for pinned_url in pinned_urls:
                    request = self._client.build_request(
                        "GET",
                        pinned_url,
                        headers=headers,
                        extensions=extensions,
                        timeout=None,
                    )
                    try:
                        response = await self._client.send(
                            request,
                            stream=True,
                            follow_redirects=False,
                        )
                    except (httpx.ConnectError, httpx.ConnectTimeout) as error:
                        last_connect_error = error
                        continue
                    try:
                        if response.is_redirect:
                            location = response.headers.get("location")
                            if location is None:
                                response.raise_for_status()
                            redirect_url = current_url.join(location)
                            break
                        return await self._read_response(response, current_url)
                    finally:
                        await response.aclose()
                if redirect_url is not None:
                    current_url = redirect_url
                    continue
                if last_connect_error is not None:
                    raise last_connect_error
                raise RuntimeError("HTTP request did not produce a response")
            raise ValueError("too many HTTP redirects")

    async def _read_response(
        self, response: httpx.Response, original_url: httpx.URL
    ) -> dict[str, Any]:
        response.raise_for_status()
        content = bytearray()
        truncated = False
        async for chunk in response.aiter_bytes():
            remaining = self._max_bytes - len(content)
            if len(chunk) > remaining:
                content.extend(chunk[:remaining])
                truncated = True
                break
            content.extend(chunk)
        media_type = response.headers.get("content-type", "")
        text = bytes(content).decode(response.encoding or "utf-8", errors="replace")
        if "html" in media_type.lower():
            try:
                soup = BeautifulSoup(text, "html.parser")
                for element in soup(["script", "style", "noscript"]):
                    element.decompose()
                text = " ".join(soup.get_text(" ", strip=True).split())
            except (AttributeError, TypeError, ValueError):
                raise WebpageError("parse_failure", retryable=False) from None
        return {
            "url": str(original_url),
            "status_code": response.status_code,
            "content_type": media_type,
            "text": text,
            "truncated": truncated,
        }

    async def aclose(self) -> None:
        if self._owns_client:
            await self._client.aclose()


def _validate_public_http_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("URL must use HTTP(S)")
    if parsed.username or parsed.password:
        raise ValueError("URL credentials are not allowed")
    hostname = parsed.hostname.rstrip(".").lower()
    if hostname == "localhost" or hostname.endswith(".localhost"):
        raise ValueError("URL must refer to a public host")
    try:
        address = ipaddress.ip_address(hostname)
    except ValueError:
        return
    if not address.is_global:
        raise ValueError("URL must refer to a public host")


async def _pin_public_url(
    url: httpx.URL,
) -> tuple[tuple[httpx.URL, ...], dict[str, str], dict[str, Any]]:
    _validate_public_http_url(str(url))
    hostname = url.host.rstrip(".").lower()
    port = url.port or (443 if url.scheme == "https" else 80)
    try:
        literal = ipaddress.ip_address(hostname)
    except ValueError:
        records = await asyncio.to_thread(
            socket.getaddrinfo,
            hostname,
            port,
            type=socket.SOCK_STREAM,
        )
        addresses = list(dict.fromkeys(record[4][0] for record in records))
    else:
        addresses = [str(literal)]
    if not addresses:
        raise ValueError("URL hostname did not resolve")
    parsed_addresses = [ipaddress.ip_address(address) for address in addresses]
    if any(not _is_global_address(address) for address in parsed_addresses):
        raise ValueError("URL must refer to a public host")
    pinned = tuple(url.copy_with(host=str(address)) for address in parsed_addresses)
    host_header = url.netloc.decode("ascii")
    headers = {"Host": host_header, "Connection": "close"}
    extensions: dict[str, Any] = {}
    if url.scheme == "https":
        extensions["sni_hostname"] = hostname.encode("idna").decode("ascii")
    return pinned, headers, extensions


def _is_global_address(address: ipaddress.IPv4Address | ipaddress.IPv6Address) -> bool:
    return (
        address.is_global
        and not address.is_private
        and not address.is_loopback
        and not address.is_link_local
        and not address.is_multicast
        and not address.is_reserved
        and not address.is_unspecified
    )


class AttachmentReader:
    def __init__(self, task_root: Path, registered: set[Path] | frozenset[Path]) -> None:
        self.task_root = task_root.resolve()
        resolved: set[Path] = set()
        registered_file_identities: dict[Path, tuple[int, int]] = {}
        for path in registered:
            candidate = path if path.is_absolute() else self.task_root / path
            candidate = candidate.resolve()
            if not _is_relative_to(candidate, self.task_root):
                raise ValueError(f"attachment is outside task root: {path}")
            resolved.add(candidate)
            try:
                identity = os.stat(candidate, follow_symlinks=False)
            except OSError:
                continue
            registered_file_identities[candidate] = (identity.st_dev, identity.st_ino)
        self.registered = frozenset(resolved)
        self._registered_file_identities = registered_file_identities

    def resolve_registered(self, path: str) -> Path:
        candidate = (self.task_root / path).resolve()
        if not _is_relative_to(candidate, self.task_root):
            raise ValueError(f"attachment path is outside task root: {path}")
        if candidate not in self.registered:
            raise ValueError(f"attachment is not registered for this task: {path}")
        if not candidate.is_file():
            raise ValueError(f"attachment does not exist or is not a file: {path}")
        if candidate.stat().st_size > MAX_ATTACHMENT_BYTES:
            raise ValueError(f"attachment exceeds size limit: {path}")
        return candidate

    def read(self, path: str) -> dict[str, Any]:
        resolved = self.resolve_registered(path)
        relative = str(resolved.relative_to(self.task_root))
        suffix = resolved.suffix.lower()
        if suffix in {".txt", ".md"}:
            text = _read_bounded_utf8(resolved)
            return {"path": relative, "kind": "text", "text": text}
        if suffix in {".csv", ".tsv"}:
            _require_size_within(resolved.stat().st_size, MAX_DECODED_TEXT_BYTES, "decoded bytes")
            rows: list[list[str]] = []
            cell_count = 0
            with resolved.open(encoding="utf-8", newline="") as stream:
                for row_number, row in enumerate(
                    csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ","),
                    start=1,
                ):
                    _require_size_within(row_number, MAX_TABULAR_ROWS, "CSV rows")
                    cell_count += len(row)
                    _require_size_within(cell_count, MAX_TABULAR_CELLS, "CSV cells")
                    rows.append(row)
            return {"path": relative, "kind": "csv", "rows": rows}
        if suffix == ".json":
            return {
                "path": relative,
                "kind": "json",
                "data": json.loads(_read_bounded_utf8(resolved)),
            }
        if suffix == ".pdf":
            with _bounded_pdf_decoding():
                reader = PdfReader(resolved)
                _require_size_within(len(reader.pages), MAX_PDF_PAGES, "PDF pages")
                pages: list[str] = []
                text_bytes = 0
                for page in reader.pages:
                    page.get_contents()
                    text = page.extract_text() or ""
                    text_bytes += len(text.encode("utf-8"))
                    _require_size_within(text_bytes, MAX_PDF_TEXT_BYTES, "PDF text bytes")
                    pages.append(text)
            return {
                "path": relative,
                "kind": "pdf",
                "page_count": len(pages),
                "text": "\n\n".join(pages),
            }
        if suffix == ".xlsx":
            _validate_xlsx_archive(resolved)
            workbook = load_workbook(resolved, read_only=True, data_only=True)
            sheets: dict[str, list[list[Any]]] = {}
            row_count = 0
            cell_count = 0
            try:
                _require_size_within(
                    len(workbook.worksheets), MAX_XLSX_WORKSHEETS, "XLSX worksheets"
                )
                for sheet in workbook.worksheets:
                    rows: list[list[Any]] = []
                    for row in sheet.iter_rows(values_only=True):
                        row_count += 1
                        _require_size_within(row_count, MAX_TABULAR_ROWS, "XLSX rows")
                        cell_count += len(row)
                        _require_size_within(cell_count, MAX_TABULAR_CELLS, "XLSX cells")
                        rows.append([_jsonable(value) for value in row])
                    sheets[sheet.title] = rows
            finally:
                workbook.close()
            return {"path": relative, "kind": "xlsx", "sheets": sheets}
        media_type = mimetypes.guess_type(resolved.name)[0] or "application/octet-stream"
        if media_type.startswith("image/"):
            kind = "image"
        elif media_type.startswith("audio/"):
            kind = "audio"
        else:
            kind = "binary"
        return {
            "path": relative,
            "kind": kind,
            "media_type": media_type,
            "size_bytes": resolved.stat().st_size,
        }


class _NamedBytesIO(io.BytesIO):
    def __init__(self, content: bytes, name: str) -> None:
        super().__init__(content)
        self.name = name


def _audio_file_state(value: os.stat_result) -> tuple[int, int, int, int, int]:
    return (
        value.st_dev,
        value.st_ino,
        value.st_size,
        value.st_mtime_ns,
        value.st_ctime_ns,
    )


def _require_audio_stat(value: os.stat_result) -> None:
    if not stat.S_ISREG(value.st_mode):
        raise ValueError("attachment could not be safely read")
    if value.st_size > MAX_AUDIO_ATTACHMENT_BYTES:
        raise ValueError("attachment exceeds audio size limit")


def _snapshot_registered_audio(reader: AttachmentReader, path: str) -> _NamedBytesIO:
    descriptor: int | None = None
    try:
        resolved = reader.resolve_registered(path)
        if resolved.suffix.lower() not in _AUDIO_TRANSCRIPTION_SUFFIXES:
            raise ValueError("attachment is not a supported audio format")
        before = os.stat(resolved, follow_symlinks=False)
        _require_audio_stat(before)
        descriptor = os.open(resolved, os.O_RDONLY | os.O_NOFOLLOW)
        opened = os.fstat(descriptor)
        _require_audio_stat(opened)
        registered_identity = reader._registered_file_identities.get(resolved)
        if registered_identity != (opened.st_dev, opened.st_ino):
            raise ValueError("attachment could not be safely read")
        if (before.st_dev, before.st_ino) != (opened.st_dev, opened.st_ino):
            raise ValueError("attachment could not be safely read")

        content = bytearray()
        while len(content) <= MAX_AUDIO_ATTACHMENT_BYTES:
            chunk = os.read(
                descriptor,
                min(1024 * 1024, MAX_AUDIO_ATTACHMENT_BYTES + 1 - len(content)),
            )
            if not chunk:
                break
            content.extend(chunk)
        if len(content) > MAX_AUDIO_ATTACHMENT_BYTES:
            raise ValueError("attachment exceeds audio size limit")

        after = os.fstat(descriptor)
        _require_audio_stat(after)
        if _audio_file_state(opened) != _audio_file_state(after) or len(content) != after.st_size:
            raise ValueError("attachment could not be safely read")
        return _NamedBytesIO(bytes(content), resolved.name)
    except OSError:
        raise ValueError("attachment could not be safely read") from None
    finally:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                raise ValueError("attachment could not be safely read") from None


class AttachmentResourceLimitError(ValueError):
    pass


def _require_size_within(actual: int, limit: int, resource: str) -> None:
    if actual > limit:
        raise AttachmentResourceLimitError(
            f"attachment resource limit exceeded for {resource}: {actual} > {limit}"
        )


def _read_bounded_utf8(path: Path) -> str:
    size = path.stat().st_size
    _require_size_within(size, MAX_DECODED_TEXT_BYTES, "decoded bytes")
    return path.read_bytes().decode("utf-8")


def _validate_xlsx_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        _require_size_within(len(entries), MAX_XLSX_ARCHIVE_ENTRIES, "XLSX archive entries")
        expanded_bytes = 0
        for entry in entries:
            if entry.flag_bits & 1:
                raise ValueError("encrypted XLSX attachments are not supported")
            expanded_bytes += entry.file_size
            _require_size_within(
                expanded_bytes,
                MAX_XLSX_UNCOMPRESSED_BYTES,
                "XLSX expanded bytes",
            )


@contextmanager
def _bounded_pdf_decoding() -> Iterator[None]:
    with _PDF_FILTER_LOCK:
        original_limits = {
            name: getattr(pypdf_filters, name) for name in _PDF_FILTER_LIMIT_NAMES
        }
        original_decompress = pypdf_filters.decompress
        for name, current_limit in original_limits.items():
            setattr(
                pypdf_filters,
                name,
                min(current_limit, MAX_PDF_DECODED_STREAM_BYTES),
            )
        pypdf_filters.decompress = _bounded_zlib_decompress
        try:
            yield
        except LimitReachedError:
            raise AttachmentResourceLimitError(
                "attachment resource limit exceeded for PDF decoded stream bytes"
            ) from None
        finally:
            pypdf_filters.decompress = original_decompress
            for name, original_limit in original_limits.items():
                setattr(pypdf_filters, name, original_limit)


def _bounded_zlib_decompress(data: bytes) -> bytes:
    decompressor = zlib.decompressobj()
    result = decompressor.decompress(data, MAX_PDF_DECODED_STREAM_BYTES + 1)
    if (
        len(result) > MAX_PDF_DECODED_STREAM_BYTES
        or decompressor.unconsumed_tail
        or not decompressor.eof
    ):
        raise LimitReachedError("PDF decoded stream exceeds resource limit")
    result += decompressor.flush(MAX_PDF_DECODED_STREAM_BYTES + 1 - len(result))
    if len(result) > MAX_PDF_DECODED_STREAM_BYTES:
        raise LimitReachedError("PDF decoded stream exceeds resource limit")
    return result


class HostedOpenAIBackend:
    provider = "openai"
    def __init__(
        self,
        *,
        client: Any,
        model: str,
        reasoning_effort: str,
        attachment_reader: AttachmentReader,
        timeout_seconds: float,
        web_search_timeout_seconds: float | None = None,
    ) -> None:
        self._client = client
        self._model = model
        self._reasoning_effort = reasoning_effort
        self._attachments = attachment_reader
        self._timeout_seconds = timeout_seconds
        self._web_search_timeout_seconds = min(
            timeout_seconds if web_search_timeout_seconds is None else web_search_timeout_seconds,
            MAX_HOSTED_WEB_SEARCH_TIMEOUT_SECONDS,
        )
        self._remote_file_ids: dict[Path, str] = {}

    async def web_search(self, query: str) -> dict[str, Any]:
        response = await asyncio.wait_for(
            self._client.responses.create(
                model=self._model,
                input=query,
                tools=[{"type": "web_search"}],
                tool_choice="required",
                include=["web_search_call.action.sources"],
                reasoning={"effort": self._reasoning_effort},
            ),
            timeout=self._web_search_timeout_seconds,
        )
        sources: list[dict[str, Any]] = []
        for item in response.output:
            if _field(item, "type") != "web_search_call":
                continue
            action = _field(item, "action")
            for source in _field(action, "sources", []) or []:
                url = _field(source, "url")
                if url:
                    sources.append({"url": url, "title": _field(source, "title")})
        result = {"answer": response.output_text, "sources": sources}
        return _with_response_usage(result, response)

    async def search(self, query: str) -> dict[str, Any]:
        return await self.web_search(query)

    async def execute_python(self, instruction: str, paths: list[str]) -> dict[str, Any]:
        file_ids = [await self._upload_registered(path) for path in paths]
        response = await asyncio.wait_for(
            self._client.responses.create(
                model=self._model,
                input=instruction,
                tools=[
                    {
                        "type": "code_interpreter",
                        "container": {
                            "type": "auto",
                            "memory_limit": "1g",
                            "file_ids": file_ids,
                        },
                    }
                ],
                tool_choice="required",
                include=["code_interpreter_call.outputs"],
                reasoning={"effort": self._reasoning_effort},
            ),
            timeout=self._timeout_seconds,
        )
        outputs: list[Any] = []
        for item in response.output:
            if _field(item, "type") == "code_interpreter_call":
                outputs.extend(_jsonable(output) for output in (_field(item, "outputs", []) or []))
        result = {"output_text": response.output_text, "outputs": outputs}
        return _with_response_usage(result, response)

    async def transcribe_audio(self, path: str) -> dict[str, Any]:
        async with asyncio.timeout(self._timeout_seconds):
            stream = await asyncio.to_thread(_snapshot_registered_audio, self._attachments, path)
            with stream:
                response = await self._client.audio.transcriptions.create(
                    model="gpt-4o-mini-transcribe",
                    file=stream,
                )
        return _with_response_usage({"text": response.text}, response)

    async def _upload_registered(self, path: str) -> str:
        resolved = self._attachments.resolve_registered(path)
        if resolved.suffix.lower() in _AUDIO_TRANSCRIPTION_SUFFIXES:
            raise ValueError("audio attachments must use transcribe_audio")
        if remote_id := self._remote_file_ids.get(resolved):
            return remote_id
        with resolved.open("rb") as stream:
            uploaded = await asyncio.wait_for(
                self._client.files.create(file=stream, purpose="assistants"),
                timeout=self._timeout_seconds,
            )
        self._remote_file_ids[resolved] = uploaded.id
        return uploaded.id


def _field(value: Any, name: str, default: Any = None) -> Any:
    if isinstance(value, dict):
        return value.get(name, default)
    return getattr(value, name, default)


def _with_response_usage(result: dict[str, Any], response: Any) -> dict[str, Any]:
    raw_usage = _field(response, "usage")
    if raw_usage is None:
        return result
    output_details = _field(raw_usage, "output_tokens_details")
    usage = _validated_hosted_usage(
        {
            "input_tokens": _field(raw_usage, "input_tokens", 0),
            "output_tokens": _field(raw_usage, "output_tokens", 0),
            "reasoning_tokens": _field(output_details, "reasoning_tokens", 0),
            "total_tokens": _field(raw_usage, "total_tokens", 0),
        }
    )
    return _TrustedHostedResult({
        **result,
        HOSTED_USAGE_METADATA_KEY: usage,
    })


def _validated_hosted_usage(value: Any) -> dict[str, int]:
    expected = set(Usage.model_fields)
    if not isinstance(value, dict) or set(value) != expected:
        raise ValueError("invalid hosted usage metadata")
    try:
        usage = Usage.model_validate(value, strict=True)
    except (TypeError, ValueError):
        raise ValueError("invalid hosted usage metadata") from None
    if any(getattr(usage, field) < 0 for field in expected):
        raise ValueError("invalid hosted usage metadata")
    return usage.model_dump(mode="json")


def _jsonable(value: Any) -> Any:
    if isinstance(value, (dict, list, str, int, float, bool)) or value is None:
        return value
    if hasattr(value, "model_dump"):
        return value.model_dump(mode="json")
    return str(value)


class FakeHostedBackend:
    provider = "openai"
    def __init__(
        self,
        api_key: str,
        hf_token: str,
        attachment_reader: AttachmentReader,
    ) -> None:
        self._api_key = api_key
        self._hf_token = hf_token
        self._attachments = attachment_reader

    async def web_search(self, query: str) -> dict[str, Any]:
        if query == "__GAIA_TIMEOUT__":
            await asyncio.sleep(60)
        if query == "__GAIA_SECRET_ERROR__":
            raise RuntimeError(f"backend rejected {self._api_key} and {self._hf_token}")
        if query == "__GAIA_NESTED_SECRET__":
            return {
                "answer": "safe",
                "sources": [],
                "debug": {
                    "api_key": self._api_key,
                    "message": f"{self._api_key} and {self._hf_token}",
                },
            }
        return {
            "answer": f"deterministic search result for {query}",
            "sources": [{"url": "https://example.test/source", "title": "Example"}],
        }

    async def search(self, query: str) -> dict[str, Any]:
        return await self.web_search(query)

    async def execute_python(self, instruction: str, paths: list[str]) -> dict[str, Any]:
        if instruction == "__GAIA_LARGE_OUTPUT__":
            return {"output_text": "x" * (MAX_TOOL_OUTPUT_BYTES + 10_000), "outputs": []}
        if instruction == "__GAIA_ESCAPED_OUTPUT__":
            return {"output_text": '"\\' * MAX_TOOL_OUTPUT_BYTES, "outputs": []}
        return {"output_text": f"deterministic code result for {instruction}", "outputs": []}

    async def transcribe_audio(self, path: str) -> dict[str, Any]:
        with _snapshot_registered_audio(self._attachments, path):
            pass
        return {"text": "deterministic audio transcript"}


@dataclass(frozen=True)
class MCPServerSettings:
    task_root: Path
    attachment_reader: AttachmentReader
    web_reader: WebReader
    backend: HostedBackend
    api_key: str
    hf_token: str
    serper_api_key: str = ""
    serpapi_api_key: str = ""
    search_provider: SearchProvider | None = None
    search_client: httpx.AsyncClient | None = None
    webpage_retry_policy: RetryPolicy = field(
        default_factory=lambda: RetryPolicy(max_attempts=3, base_delay_seconds=0.25, max_delay_seconds=1)
    )

    @property
    def secrets(self) -> tuple[str, ...]:
        return (self.api_key, self.hf_token, self.serper_api_key, self.serpapi_api_key)

    async def aclose(self) -> None:
        close_web_reader = getattr(self.web_reader, "aclose", None)
        if close_web_reader is not None:
            await close_web_reader()
        if self.search_client is not None:
            await self.search_client.aclose()

    @classmethod
    def for_fake_backend(
        cls,
        *,
        task_root: Path,
        attachments: set[Path],
        api_key: str,
        hf_token: str,
        model: str,
        reasoning_effort: str,
        timeout_seconds: float,
    ) -> MCPServerSettings:
        attachment_reader = AttachmentReader(task_root, attachments)
        return cls(
            task_root=task_root.resolve(),
            attachment_reader=attachment_reader,
            web_reader=HTTPWebReader(
                timeout_seconds=min(timeout_seconds, MAX_DIRECT_WEBPAGE_TIMEOUT_SECONDS)
            ),
            backend=FakeHostedBackend(api_key, hf_token, attachment_reader),
            api_key=api_key,
            hf_token=hf_token,
        )


def _safe_trace_metadata(values: dict[str, Any]) -> dict[str, Any]:
    safe: dict[str, Any] = {}
    provider = values.get("provider")
    if provider in {"openai", "serper", "serpapi"}:
        safe["provider"] = provider
    category = values.get("category")
    if category in {
        "tool_error", "search_unavailable", "timeout", "transport", "http_status",
        "parse_failure", "empty_content", "policy_rejection", "provider_error",
    }:
        safe["category"] = category
    retryable = values.get("retryable")
    if isinstance(retryable, bool):
        safe["retryable"] = retryable
    for key, minimum in (("attempt_count", 1), ("circuit_open_count", 0)):
        value = values.get(key)
        if (
            isinstance(value, int)
            and not isinstance(value, bool)
            and minimum <= value <= _MAX_TRACE_COUNT
        ):
            safe[key] = value
    http_status = values.get("http_status")
    if (
        isinstance(http_status, int)
        and not isinstance(http_status, bool)
        and 100 <= http_status <= 599
    ):
        safe["http_status"] = http_status
    for key in ("fallback_attempted", "fallback_used"):
        value = values.get(key)
        if isinstance(value, bool):
            safe[key] = value
    return safe


def _result_trace_metadata(result: dict[str, Any]) -> dict[str, Any]:
    if isinstance(result, SearchProviderResult):
        return _safe_trace_metadata(
            {
                "provider": result.provider,
                "attempt_count": result.attempt_count,
                "circuit_open_count": result.circuit_open_count,
                "fallback_used": result.fallback_used,
            }
        )
    if isinstance(result, _SearchFallbackResult):
        return _safe_trace_metadata(result.trace_metadata)
    return {}


def _failure_trace_metadata(error: ToolCallError, failure: SafeToolFailure) -> dict[str, Any]:
    return _safe_trace_metadata(
        {
            **failure.model_dump(exclude={"tool"}),
            "circuit_open_count": getattr(error, "circuit_open_count", None),
            "fallback_attempted": getattr(error, "fallback_attempted", None),
        }
    )


def _call_tool_result(
    payload: dict[str, Any] | None,
    *,
    trace_metadata: dict[str, Any],
    error_message: str | None = None,
) -> CallToolResult:
    text = error_message if error_message is not None else serialize_tool_output(payload or {})
    return CallToolResult(
        content=[TextContent(type="text", text=text)],
        structuredContent=payload or {},
        isError=error_message is not None,
        _meta={_TRACE_META_KEY: trace_metadata},
    )


async def _safe_tool_call(
    settings: MCPServerSettings,
    operation: Awaitable[dict[str, Any]] | Callable[[], dict[str, Any]],
    *,
    web_result: bool = False,
    mcp_result: bool = False,
) -> dict[str, Any] | CallToolResult:
    try:
        result = await operation if hasattr(operation, "__await__") else operation()
        is_search_result = isinstance(result, _SearchFallbackResult)
        payload = dict(result)
        raw_usage = payload.pop(HOSTED_USAGE_METADATA_KEY, None)
        usage = (
            _validated_hosted_usage(raw_usage)
            if isinstance(result, _TrustedHostedResult) and raw_usage is not None
            else None
        )
        redacted = _redact(payload, settings.secrets)
        if is_search_result:
            redacted = {
                "url": str(redacted.get("url") or "")[:_MAX_FALLBACK_URL_CHARS],
                "answer": str(redacted.get("answer") or ""),
                "sources": redacted.get("sources")
                if isinstance(redacted.get("sources"), list)
                else [],
                "fallback": "search",
            }
        metadata = {HOSTED_USAGE_METADATA_KEY: usage} if usage is not None else None
        bounded = (
            _bounded_web_search_result(redacted, metadata=metadata)
            if web_result or is_search_result
            else _bounded_result(redacted, metadata=metadata)
        )
        trace_metadata = _result_trace_metadata(result)
        if mcp_result and trace_metadata:
            return _call_tool_result(bounded, trace_metadata=trace_metadata)
        return bounded
    except ToolCallError as error:
        failure = SafeToolFailure.from_exception("tool", error)
        details = failure.model_dump(exclude={"tool", "http_status"})
        if failure.http_status is not None:
            details["http_status"] = failure.http_status
        message = serialize_tool_output({"error": "tool_failed", **details})
        if mcp_result:
            return _call_tool_result(
                None,
                trace_metadata=_failure_trace_metadata(error, failure),
                error_message=message,
            )
        raise RuntimeError(message) from None
    # This is the server boundary: redact arbitrary backend/parser failures before
    # MCP serializes them into consumer-visible tool errors.
    except Exception as error:  # noqa: BLE001
        message = _replace_secrets(str(error), settings.secrets)
        raise RuntimeError(message) from None


async def _read_webpage_with_fallback(settings: MCPServerSettings, url: str) -> dict[str, Any]:
    for attempt in range(1, settings.webpage_retry_policy.max_attempts + 1):
        try:
            result = await settings.web_reader.read(url)
        except ValueError as error:
            category = "policy_rejection" if _is_webpage_policy_error(error) else "parse_failure"
            raise WebpageError(category, retryable=False, attempt_count=attempt) from None
        except ToolCallError as error:
            error.attempt_count = attempt
            if not bool(getattr(error, "retryable", False)):
                raise
            if attempt < settings.webpage_retry_policy.max_attempts:
                await asyncio.sleep(retry_delay(settings.webpage_retry_policy, attempt, lambda: 0.0))
                continue
            try:
                provider = settings.search_provider or settings.backend
                search = provider.search if hasattr(provider, "search") else provider.web_search
                snippets = await search(url)
            except ToolCallError:
                raise error from None
            provider_name = getattr(snippets, "provider", getattr(provider, "provider", None))
            trace_metadata: dict[str, Any] = {
                "attempt_count": getattr(snippets, "attempt_count", 1),
                "fallback_used": True,
            }
            if provider_name in {"openai", "serper", "serpapi"}:
                trace_metadata["provider"] = provider_name
            circuit_open_count = getattr(snippets, "circuit_open_count", None)
            if isinstance(circuit_open_count, int) and not isinstance(circuit_open_count, bool):
                trace_metadata["circuit_open_count"] = circuit_open_count
            return _SearchFallbackResult(
                {
                "url": url,
                "answer": str(snippets.get("answer") or ""),
                "sources": snippets.get("sources") if isinstance(snippets.get("sources"), list) else [],
                "fallback": "search",
                },
                trace_metadata,
            )
        if not str(result.get("text") or "").strip():
            raise WebpageError("empty_content", retryable=False, attempt_count=attempt)
        return result
    raise AssertionError("bounded webpage retry did not return or raise")


def _is_webpage_policy_error(error: ValueError) -> bool:
    return str(error).startswith(
        (
            "URL must use HTTP(S)",
            "URL credentials are not allowed",
            "URL must refer to a public host",
            "URL hostname did not resolve",
        )
    )


def build_server(settings: MCPServerSettings) -> MCPServer:
    server = MCPServer("gaia-tools")

    @server.tool()
    async def web_search(query: str) -> dict[str, Any]:
        """Search the public and return a concise answer with source URLs."""
        return await _safe_tool_call(
            settings,
            (
                (settings.search_provider or settings.backend).search(query)
                if hasattr(settings.search_provider or settings.backend, "search")
                else settings.backend.web_search(query)
            ),
            web_result=True,
            mcp_result=True,
        )

    @server.tool()
    async def read_webpage(url: str) -> dict[str, Any]:
        """Read public HTTP(S) page text with redirects, timeout, and size limits."""
        return await _safe_tool_call(
            settings, _read_webpage_with_fallback(settings, url), mcp_result=True
        )

    @server.tool()
    async def execute_python(instruction: str, paths: list[str] | None = None) -> dict[str, Any]:
        """Run a Python analysis instruction in hosted OpenAI Code Interpreter."""
        return await _safe_tool_call(
            settings, settings.backend.execute_python(instruction, paths or []), mcp_result=True
        )

    @server.tool()
    async def read_attachment(path: str) -> dict[str, Any]:
        """Extract text and metadata from an attachment registered for this task."""
        return await _safe_tool_call(
            settings, lambda: settings.attachment_reader.read(path), mcp_result=True
        )

    @server.tool()
    async def transcribe_audio(path: str) -> dict[str, Any]:
        """Transcribe a registered audio attachment with hosted speech recognition."""
        return await _safe_tool_call(
            settings, settings.backend.transcribe_audio(path), mcp_result=True
        )

    @server.tool()
    async def submit_answer(answer: str, evidence_ids: list[str]) -> dict[str, Any]:
        """Return the concise terminal answer and evidence references."""
        return await _safe_tool_call(
            settings,
            lambda: {"answer": answer.strip(), "evidence_ids": evidence_ids},
            mcp_result=True,
        )

    return server


def _settings_from_env() -> MCPServerSettings:
    task_root = Path(os.environ["GAIA_TASK_ROOT"]).resolve()
    raw_attachments = json.loads(os.environ.get("GAIA_ATTACHMENTS", "[]"))
    if not isinstance(raw_attachments, list) or not all(
        isinstance(path, str) for path in raw_attachments
    ):
        raise ValueError("GAIA_ATTACHMENTS must be a JSON list of paths")
    attachment_reader = AttachmentReader(task_root, {Path(path) for path in raw_attachments})
    api_key = os.environ.get("OPENAI_API_KEY", "")
    hf_token = os.environ.get("HF_TOKEN", "")
    serper_api_key = os.environ.get("SERPER_API_KEY", "")
    serpapi_api_key = os.environ.get("SERPAPI_API_KEY", "")
    timeout_seconds = float(os.environ.get("GAIA_TIMEOUT_SECONDS", "300"))
    provider_retry_policy = RetryPolicy(
        max_attempts=int(os.environ.get("GAIA_PROVIDER_MAX_ATTEMPTS", "4")),
        base_delay_seconds=float(os.environ.get("GAIA_PROVIDER_BACKOFF_BASE", "0.5")),
    )
    search_circuit_cooldown = float(os.environ.get("GAIA_SEARCH_CIRCUIT_COOLDOWN", "30"))
    fake_backend = os.environ.get("GAIA_MCP_FAKE_BACKEND") == "1"
    fallback_provider_count = int(bool(serper_api_key)) + int(bool(serpapi_api_key))
    hosted_web_search_timeout_seconds = min(
        timeout_seconds / (fallback_provider_count + 1),
        MAX_HOSTED_WEB_SEARCH_TIMEOUT_SECONDS,
    )
    if fake_backend:
        backend: HostedBackend = FakeHostedBackend(api_key, hf_token, attachment_reader)
    else:
        client = AsyncOpenAI(api_key=api_key, max_retries=0)
        backend = HostedOpenAIBackend(
            client=client,
            model=os.environ.get("GAIA_MODEL", "gpt-5.6-luna"),
            reasoning_effort=os.environ.get("GAIA_REASONING_EFFORT", "max"),
            attachment_reader=attachment_reader,
            timeout_seconds=timeout_seconds,
            web_search_timeout_seconds=hosted_web_search_timeout_seconds,
        )
    search_provider: SearchProvider | None = None
    search_client: httpx.AsyncClient | None = None
    if not fake_backend:
        search_providers: list[SearchProvider] = [OpenAISearchProvider(backend)]
        search_client = httpx.AsyncClient()
        if serper_api_key:
            search_providers.append(
                SerperSearchProvider(
                    serper_api_key, search_client, timeout_seconds=min(timeout_seconds, 20.0)
                )
            )
        if serpapi_api_key:
            search_providers.append(
                SerpAPISearchProvider(
                    serpapi_api_key, search_client, timeout_seconds=min(timeout_seconds, 20.0)
                )
            )
        search_provider = SearchProviderChain(
            search_providers,
            provider_retry_policy,
            time.monotonic,
            circuit_cooldown_seconds=search_circuit_cooldown,
            operation_timeout_seconds=timeout_seconds,
        )
    return MCPServerSettings(
        task_root=task_root,
        attachment_reader=attachment_reader,
        web_reader=HTTPWebReader(
            timeout_seconds=min(timeout_seconds, MAX_DIRECT_WEBPAGE_TIMEOUT_SECONDS)
        ),
        backend=backend,
        api_key=api_key,
        hf_token=hf_token,
        serper_api_key=serper_api_key,
        serpapi_api_key=serpapi_api_key,
        search_provider=search_provider,
        search_client=search_client,
    )


if __name__ == "__main__":
    _settings = _settings_from_env()
    try:
        build_server(_settings).run()
    finally:
        asyncio.run(_settings.aclose())
